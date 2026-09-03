from __future__ import annotations

import csv
import hashlib
import io
from pathlib import Path
from typing import Iterable
from zipfile import ZipFile, is_zipfile

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph as WordParagraph
from openpyxl import load_workbook
from pptx import Presentation

from .models import ExtractedRecord, SourceMetadata


SUPPORTED_SUFFIXES = {".xlsx", ".xls", ".csv", ".pptx", ".ppt", ".docx", ".doc"}
_OOXML_ZIP_ROOTS = {".xlsx": "xl/", ".pptx": "ppt/", ".docx": "word/"}
_LEGACY_CONVERSION_TARGET = {".xls": ".xlsx", ".ppt": ".pptx", ".doc": ".docx"}


class DocumentParseError(ValueError):
    """Raised when a user-supplied document cannot be safely parsed."""


def document_id(file_name: str, payload: bytes) -> str:
    digest = hashlib.sha256(payload).hexdigest()[:16]
    return f"{Path(file_name).stem}-{digest}"


def parse_document(file_name: str, payload: bytes) -> list[ExtractedRecord]:
    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentParseError(f"Unsupported file type '{suffix or 'unknown'}'.")
    if not payload:
        raise DocumentParseError("The uploaded file is empty.")
    if suffix in _OOXML_ZIP_ROOTS and not is_zipfile(io.BytesIO(payload)):
        raise DocumentParseError(f"'{file_name}' is not a valid {suffix} file.")
    if suffix in _OOXML_ZIP_ROOTS:
        with ZipFile(io.BytesIO(payload)) as archive:
            names = set(archive.namelist())
        expected_root = _OOXML_ZIP_ROOTS[suffix]
        if "[Content_Types].xml" not in names or not any(name.startswith(expected_root) for name in names):
            raise DocumentParseError(f"'{file_name}' does not contain a valid {suffix[1:].upper()} package.")
    doc_id = document_id(file_name, payload)
    try:
        if suffix == ".csv":
            return parse_csv(file_name, payload, doc_id)
        if suffix == ".xlsx":
            return parse_excel(file_name, payload, doc_id)
        if suffix == ".pptx":
            return parse_powerpoint(file_name, payload, doc_id)
        if suffix == ".docx":
            return parse_word(file_name, payload, doc_id)
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"Could not read '{file_name}': {exc}") from exc
    raise DocumentParseError(
        f"'{suffix}' is a legacy binary format. Convert it to "
        f"'{_LEGACY_CONVERSION_TARGET[suffix]}' and upload it again."
    )


def _stringify(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _table_record(
    file_name: str,
    file_type: str,
    doc_id: str,
    headers: list[str],
    rows: Iterable[tuple[int, list[str]]],
    sheet_name: str | None = None,
    warnings: tuple[str, ...] = (),
) -> list[ExtractedRecord]:
    records: list[ExtractedRecord] = []
    labelled_headers = [header or f"Column {index + 1}" for index, header in enumerate(headers)]
    for row_number, values in rows:
        fields = [f"{header}: {value}" for header, value in zip(labelled_headers, values) if value]
        if fields:
            text_prefix = f"Sheet: {sheet_name}\n" if sheet_name else ""
            records.append(
                ExtractedRecord(
                    text=text_prefix + " | ".join(fields),
                    source=SourceMetadata(
                        file_name=file_name,
                        file_type=file_type,
                        document_id=doc_id,
                        sheet_name=sheet_name,
                        row_start=row_number,
                        row_end=row_number,
                    ),
                    warnings=warnings,
                )
            )
    return records


def parse_csv(file_name: str, payload: bytes, doc_id: str) -> list[ExtractedRecord]:
    decoded = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            decoded = payload.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise DocumentParseError("Could not determine the CSV text encoding.")
    sample = decoded[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(decoded), dialect, strict=True)
    try:
        headers = [_stringify(value) for value in next(reader)]
    except StopIteration:
        raise DocumentParseError("The CSV file contains no rows.")
    warnings: list[str] = []
    try:
        rows = []
        for number, row in enumerate(reader, start=2):
            if len(row) != len(headers):
                warnings.append(f"CSV row {number} has {len(row)} values; expected {len(headers)}.")
            rows.append((number, [_stringify(value) for value in row]))
    except csv.Error as exc:
        raise DocumentParseError(f"Malformed CSV row: {exc}") from exc
    return _table_record(file_name, "csv", doc_id, headers, rows, warnings=tuple(warnings))


def parse_excel(file_name: str, payload: bytes, doc_id: str) -> list[ExtractedRecord]:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    formula_workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    records: list[ExtractedRecord] = []
    for worksheet, formula_worksheet in zip(workbook.worksheets, formula_workbook.worksheets):
        headers: list[str] = []
        data_rows: list[tuple[int, list[str]]] = []
        warnings: list[str] = []
        rows = zip(worksheet.iter_rows(), formula_worksheet.iter_rows())
        for row_number, (result_row, formula_row) in enumerate(rows, start=1):
            values: list[str] = []
            for result_cell, formula_cell in zip(result_row, formula_row):
                if isinstance(formula_cell.value, str) and formula_cell.value.startswith("=") and result_cell.value is None:
                    warnings.append(f"Formula result unavailable in {worksheet.title}!{formula_cell.coordinate}")
                values.append(_stringify(result_cell.value))
            if row_number == 1:
                headers = values
            else:
                data_rows.append((row_number, values))
        if not headers and not data_rows:
            continue
        records.extend(_table_record(file_name, "excel", doc_id, headers, data_rows, worksheet.title, tuple(warnings)))
    if not records:
        raise DocumentParseError("No populated data rows were found in this workbook.")
    return records


def _word_table_records(
    file_name: str,
    doc_id: str,
    headers: list[str],
    rows: Iterable[tuple[int, list[str]]],
    section_label: str,
) -> list[ExtractedRecord]:
    records: list[ExtractedRecord] = []
    labelled_headers = [header or f"Column {index + 1}" for index, header in enumerate(headers)]
    for row_number, values in rows:
        fields = [f"{header}: {value}" for header, value in zip(labelled_headers, values) if value]
        if fields:
            records.append(
                ExtractedRecord(
                    text=f"{section_label}\n" + " | ".join(fields),
                    source=SourceMetadata(
                        file_name=file_name,
                        file_type="word",
                        document_id=doc_id,
                        section_title=section_label,
                        row_start=row_number,
                        row_end=row_number,
                    ),
                )
            )
    return records


def _iter_block_items(document: Document) -> Iterable[WordParagraph | WordTable]:
    """Yield paragraphs and tables from a Word document body, in document order."""
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield WordParagraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield WordTable(child, document)


def parse_word(file_name: str, payload: bytes, doc_id: str) -> list[ExtractedRecord]:
    document = Document(io.BytesIO(payload))
    records: list[ExtractedRecord] = []
    current_section: str | None = None
    section_buffer: list[str] = []
    section_start: int | None = None
    section_end: int | None = None
    paragraph_index = 0
    table_index = 0

    def flush_section() -> None:
        nonlocal section_buffer, section_start, section_end
        if section_buffer:
            prefix = f"Section: {current_section}\n" if current_section else ""
            records.append(
                ExtractedRecord(
                    text=prefix + "\n".join(section_buffer),
                    source=SourceMetadata(
                        file_name=file_name,
                        file_type="word",
                        document_id=doc_id,
                        section_title=current_section,
                        paragraph_start=section_start,
                        paragraph_end=section_end,
                    ),
                )
            )
        section_buffer = []
        section_start = None
        section_end = None

    for block in _iter_block_items(document):
        if isinstance(block, WordParagraph):
            paragraph_index += 1
            text = _stringify(block.text)
            if not text:
                continue
            style_name = (block.style.name or "") if block.style else ""
            if style_name.lower().startswith("heading"):
                flush_section()
                current_section = text
                continue
            if section_start is None:
                section_start = paragraph_index
            section_end = paragraph_index
            section_buffer.append(text)
        else:
            flush_section()
            table_index += 1
            label = f"{current_section} - Table {table_index}" if current_section else f"Table {table_index}"
            rows = list(block.rows)
            headers = [_stringify(cell.text) for cell in rows[0].cells] if rows else []
            data_rows = [
                (row_number, [_stringify(cell.text) for cell in row.cells])
                for row_number, row in enumerate(rows[1:], start=2)
            ]
            records.extend(_word_table_records(file_name, doc_id, headers, data_rows, label))
    flush_section()
    if not records:
        raise DocumentParseError("No readable text or tables were found in this document.")
    return records


def parse_powerpoint(file_name: str, payload: bytes, doc_id: str) -> list[ExtractedRecord]:
    presentation = Presentation(io.BytesIO(payload))
    records: list[ExtractedRecord] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        title = _stringify(slide.shapes.title.text) if slide.shapes.title else None
        fragments: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = _stringify(shape.text)
                if text and text != title:
                    fragments.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [_stringify(cell.text) for cell in row.cells]
                    if any(cells):
                        fragments.append(" | ".join(cells))
        notes = (
            _stringify(slide.notes_slide.notes_text_frame.text)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame
            else ""
        )
        if notes:
            fragments.append(f"Speaker notes: {notes}")
        text_parts = ([f"Title: {title}"] if title else []) + fragments
        if text_parts:
            records.append(
                ExtractedRecord(
                    text="\n".join(text_parts),
                    source=SourceMetadata(
                        file_name=file_name,
                        file_type="powerpoint",
                        document_id=doc_id,
                        slide_number=slide_number,
                        slide_title=title,
                    ),
                )
            )
    if not records:
        raise DocumentParseError("No readable text or tables were found in this presentation.")
    return records
