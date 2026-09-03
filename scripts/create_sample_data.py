"""Generate small Acme Retail files used in the interview demo."""

from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation


SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_data"


def main() -> None:
    SAMPLE_DIR.mkdir(exist_ok=True)
    workbook = Workbook()
    sales = workbook.active
    sales.title = "Regional Sales"
    sales.append(["Region", "Q1 Revenue", "Q1 Target", "Target Attainment"])
    sales.append(["North", 3800000, 4000000, "95%"])
    sales.append(["South", 4200000, 4000000, "105%"])
    sales.append(["East", 3100000, 3500000, "89%"])
    workbook.save(SAMPLE_DIR / "sales_q1.xlsx")

    (SAMPLE_DIR / "customer_sales.csv").write_text(
        "Segment,Region,Q1 Customers,Retention\nEnterprise,South,125,94%\nSMB,South,410,88%\nEnterprise,North,110,91%\n",
        encoding="utf-8",
    )

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Q1 Business Review"
    slide.placeholders[1].text = "South region exceeded its Q1 revenue target by 5%."
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Regional Performance"
    slide.placeholders[1].text = "South delivered $4.2M against a $4.0M target. Enterprise retention in South was 94%."
    deck.save(SAMPLE_DIR / "q1_business_review.pptx")


if __name__ == "__main__":
    main()
